"""The complete shop audit report (docs/16 §17 — docs/17 P10): the audit sections as one printable
page (the browser's "Save as PDF", which shapes Urdu correctly) or one Excel workbook with a
sheet per section. Every figure is the same service call the screens use, so the export and the
dashboard cannot disagree."""

import datetime as dt
import html
import io
import uuid
from decimal import Decimal

from fastapi.responses import Response
from sqlalchemy.orm import Session

from app.core.periods import Period
from app.modules.shopaudit import (
    finance,
    inventory,
    parties,
    products,
    reconcile,
    security,
    service,
)

SECTIONS = (
    "summary",
    "pnl",
    "cash",
    "expenses",
    "tax",
    "customers",
    "suppliers",
    "udhaar",
    "inventory",
    "ageing",
    "products",
    "invoices",
    "verification",
    "checks",
    "security",
)
TITLES = {
    "summary": "Executive summary",
    "pnl": "Profit & loss",
    "cash": "Cash & bank",
    "expenses": "Expenses",
    "tax": "Tax",
    "customers": "Top customers",
    "suppliers": "Top suppliers",
    "udhaar": "Udhaar (receivables / payables)",
    "inventory": "Inventory valuation",
    "ageing": "Stock ageing",
    "products": "Product performance",
    "invoices": "Invoice statistics",
    "verification": "ID checks",
    "checks": "Reconciliation checks",
    "security": "Security & fraud audit",
}


def _s(v) -> str:
    if v is None:
        return ""
    if isinstance(v, Decimal):
        return f"{v:,.2f}"
    if isinstance(v, dt.datetime):
        return v.strftime("%d-%m-%Y %H:%M")
    if isinstance(v, dt.date):
        return v.strftime("%d-%m-%Y")
    if isinstance(v, bool):
        return "yes" if v else "no"
    return str(v)


def _kv(rows: list[tuple]) -> tuple[list[str], list[list]]:
    return ["Item", "Value"], [[k, v] for k, v in rows]


def build(
    db: Session,
    shop_id: uuid.UUID,
    shop,
    p: Period,
    *,
    sections: list[str],
    detail: bool,
    can_cost: bool,
    is_owner: bool,
) -> list[dict]:
    """[{key, title, header, rows, notes}] in report order"""
    out: list[dict] = []
    want = [s for s in SECTIONS if s in sections]
    limit = 100 if detail else 10
    if "summary" in want:
        o = service.overview(db, shop_id, p, can_cost=can_cost)
        rows = [
            ("Net sales", o["sales"]["net_sales"]),
            ("Sale bills / phones sold", f"{o['sales']['invoices']} / {o['sales']['phones_sold']}"),
            ("Purchases (net)", o["purchases"]["purchases_net"] if can_cost else None),
            ("Expenses", o["expenses"]["total"]),
            ("Net profit", o["profit"]["net_profit"] if o["profit"] else None),
            ("Money in / out", f"{_s(o['money']['inflows'])} / {_s(o['money']['outflows'])}"),
            (
                "Cash & bank: opening → closing",
                f"{_s(o['money']['opening'])} → {_s(o['money']['closing'])}",
            ),
            ("Cash & bank now", o["money"]["current"]),
            (
                "People owe us: opening → closing",
                f"{_s(o['debt']['receivable']['opening'])} → {_s(o['debt']['receivable']['closing'])}",
            ),
            (
                "We owe people: opening → closing",
                f"{_s(o['debt']['payable']['opening'])} → {_s(o['debt']['payable']['closing'])}",
            ),
            (
                "Phones in stock: opening → closing",
                f"{o['stock']['opening']['count']} → {o['stock']['closing']['count']}",
            ),
            ("Stock value at cost (closing)", o["stock"]["closing"]["cost_value"]),
            ("Stock value at asking (closing)", o["stock"]["closing"]["asking_value"]),
            ("Tax collected", o["sales"]["tax_collected"]),
            ("Records start on", o["coverage"]["first_day"]),
        ]
        h, r = _kv(rows)
        out.append(
            {"key": "summary", "header": h, "rows": r, "notes": o["coverage"]["limitations"]}
        )
    if "pnl" in want:
        x = finance.pnl(db, shop_id, p)
        c, v = x["current"], x["previous"]
        keys = [
            ("gross_sales", "Sales before discounts"),
            ("discounts", "Discounts"),
            ("sales_returns", "Returns"),
            ("net_sales", "Net sales"),
            ("cogs", "Cost of goods sold"),
            ("gross_profit", "Gross profit"),
            ("operating_expenses", "Operating expenses"),
            ("operating_profit", "Operating profit"),
            ("repair_income", "Repair income"),
            ("other_income", "Other income"),
            ("net_profit", "Net profit"),
            ("tax_collected", "Tax collected"),
            ("owner_drawing", "Owner drawings"),
        ]
        rows = [[lbl, c.get(k), v.get(k)] for k, lbl in keys if k in c]
        out.append(
            {
                "key": "pnl",
                "header": ["Line", "This period", "Period before"],
                "rows": rows,
                "notes": x["notes"],
            }
        )
        if detail:
            out.append(
                {
                    "key": "pnl_series",
                    "title": f"Profit & loss by {x['granularity']}",
                    "header": ["From", "Net sales", "Cost", "Gross", "Expenses", "Net profit"],
                    "rows": [
                        [
                            s["bucket"],
                            s["net_sales"],
                            s["cogs"],
                            s["gross_profit"],
                            s["expenses"],
                            s["net_profit"],
                        ]
                        for s in x["series"]
                    ],
                    "notes": [],
                }
            )
    if "cash" in want:
        x = finance.cash(db, shop_id, p)
        rows = [
            [
                a["name"],
                a["kind"],
                a["opening"],
                a["cash_in"],
                a["cash_out"],
                a["transfers_in"],
                a["transfers_out"],
                a["closing"],
                a["current"],
            ]
            for a in x["accounts"]
        ]
        rows.append(
            [
                "Total",
                "",
                x["opening"],
                x["in_total"],
                x["out_total"],
                "",
                "",
                x["closing"],
                x["current"],
            ]
        )
        out.append(
            {
                "key": "cash",
                "header": [
                    "Account",
                    "Type",
                    "Opening",
                    "In",
                    "Out",
                    "Moved in",
                    "Moved out",
                    "Closing",
                    "Now",
                ],
                "rows": rows,
                "notes": x["notes"],
            }
        )
        reasons = [["In · " + k, v] for k, v in x["in"].items()] + [
            ["Out · " + k, v] for k, v in x["out"].items()
        ]
        out.append(
            {
                "key": "cash_reasons",
                "title": "Where money came from and went",
                "header": ["Reason", "Amount"],
                "rows": reasons,
                "notes": [],
            }
        )
        intl = finance.international(db, shop_id, p)
        if intl["rows"]:
            out.append(
                {
                    "key": "international",
                    "title": "Remittances from abroad",
                    "header": [
                        "Date",
                        "Bill",
                        "Person",
                        "Sender",
                        "Country",
                        "Service",
                        "Currency",
                        "Foreign amount",
                        "Rate",
                        "PKR",
                        "Reference",
                    ],
                    "rows": [
                        [
                            r["date"],
                            r["document_no"],
                            r["party_name"],
                            r["sender_name"],
                            r["country"],
                            r["service"],
                            r["currency"],
                            r["foreign_amount"],
                            r["rate"],
                            r["pkr_amount"],
                            r["reference"],
                        ]
                        for r in intl["rows"][:limit]
                    ],
                    "notes": intl["notes"],
                }
            )
    if "expenses" in want:
        x = finance.expenses(db, shop_id, p)
        rows = [
            [
                c["name"],
                c["count"],
                c["amount"],
                next(
                    (q["amount"] for q in x["previous_by_category"] if q["key"] == c["key"]), None
                ),
            ]
            for c in x["by_category"]
        ]
        rows.append(["Total (owner drawings apart)", x["count"], x["total"], x["previous_total"]])
        out.append(
            {
                "key": "expenses",
                "header": ["Category", "Entries", "This period", "Period before"],
                "rows": rows,
                "notes": x["notes"],
            }
        )
        if detail and x["largest"]:
            out.append(
                {
                    "key": "expenses_largest",
                    "title": "Largest expenses",
                    "header": ["Date", "No", "Category", "Description", "Account", "By", "Amount"],
                    "rows": [
                        [
                            e["date"],
                            e["no"],
                            e["category_name"],
                            e["description"],
                            e["account"],
                            e["user_name"],
                            e["amount"],
                        ]
                        for e in x["largest"]
                    ],
                    "notes": [],
                }
            )
    if "tax" in want:
        x = finance.tax(db, shop_id, p)
        h, r = _kv(
            [
                ("Bills", x["invoices"]),
                ("Taxed bills", x["taxed_invoices"]),
                ("Taxable sales", x["taxable_sales"]),
                ("Tax on bills", x["tax_on_invoices"]),
                ("Tax net of returns", x["tax_net_of_returns"]),
            ]
        )
        out.append({"key": "tax", "header": h, "rows": r, "notes": x["notes"]})
        if detail and x["rows"]:
            out.append(
                {
                    "key": "tax_rows",
                    "title": "Taxed bills",
                    "header": [
                        "Date",
                        "Bill",
                        "Customer",
                        "Taxable",
                        "Rate %",
                        "Incl.",
                        "Tax",
                        "Total",
                    ],
                    "rows": [
                        [
                            r["date"],
                            r["no"],
                            r["party_name"],
                            r["taxable"],
                            r["tax_rate"],
                            r["tax_inclusive"],
                            r["tax_amount"],
                            r["total"],
                        ]
                        for r in x["rows"][:limit]
                    ],
                    "notes": [],
                }
            )
    if "customers" in want:
        for metric, title in (
            ("customers_sales", "Top customers by sales"),
            ("customers_profit", "Top customers by profit"),
            ("receivable", "Largest amounts owed to us"),
        ):
            x = parties.rankings(db, shop_id, p, metric, can_cost=can_cost, limit=limit)
            out.append(
                {
                    "key": f"rank_{metric}",
                    "title": title,
                    "header": ["#", "Name", "Phone", "Value", "Count", "Share %"],
                    "rows": [
                        [r["rank"], r["name"], r["phone"], r["value"], r["count"], r["share"]]
                        for r in x["rows"]
                    ],
                    "notes": [],
                }
            )
    if "suppliers" in want and can_cost:
        for metric, title in (
            ("suppliers_purchases", "Top suppliers by purchases"),
            ("suppliers_profit", "Top suppliers by profit on their phones"),
            ("payable", "Largest amounts we owe"),
        ):
            x = parties.rankings(db, shop_id, p, metric, can_cost=can_cost, limit=limit)
            out.append(
                {
                    "key": f"rank_{metric}",
                    "title": title,
                    "header": ["#", "Name", "Phone", "Value", "Count", "Share %"],
                    "rows": [
                        [r["rank"], r["name"], r["phone"], r["value"], r["count"], r["share"]]
                        for r in x["rows"]
                    ],
                    "notes": [],
                }
            )
    if "udhaar" in want:
        x = parties.udhaar(db, shop_id, p, can_cost=can_cost, limit=limit)
        h, r = _kv(
            [
                (
                    "People owe us: opening → closing → now",
                    f"{_s(x['receivable']['opening'])} → {_s(x['receivable']['closing'])} → {_s(x['receivable']['current'])}",
                ),
                (
                    "We owe people: opening → closing → now",
                    f"{_s(x['payable']['opening'])} → {_s(x['payable']['closing'])} → {_s(x['payable']['current'])}"
                    if can_cost
                    else None,
                ),
                *[(f"Movement · {k}", v) for k, v in x["movement"].items() if v is not None],
            ]
        )
        out.append({"key": "udhaar", "header": h, "rows": r, "notes": x["notes"]})
        out.append(
            {
                "key": "udhaar_people",
                "title": "People who owe us",
                "header": ["Name", "Phone", "Amount", "Oldest bill", "Age (days)", "Last paid"],
                "rows": [
                    [
                        r["name"],
                        r["phone"],
                        r["amount"],
                        r["oldest_bill"],
                        r["age_days"],
                        r["last_payment"],
                    ]
                    for r in x["receivable_rows"]
                ],
                "notes": [],
            }
        )
        if can_cost:
            out.append(
                {
                    "key": "udhaar_we_owe",
                    "title": "People we owe",
                    "header": ["Name", "Phone", "Amount", "Oldest bill", "Age (days)", "Last paid"],
                    "rows": [
                        [
                            r["name"],
                            r["phone"],
                            r["amount"],
                            r["oldest_bill"],
                            r["age_days"],
                            r["last_payment"],
                        ]
                        for r in x["payable_rows"]
                    ],
                    "notes": [],
                }
            )
    if "inventory" in want:
        x = inventory.summary(db, shop_id, shop, p.as_of, can_cost=can_cost)
        h, r = _kv(
            [
                ("Phones in stock (as of)", x["count"]),
                ("Value at cost", x["cost_value"]),
                ("Value at asking", x["asking_value"]),
                ("Expected margin", x["expected_margin"]),
                ("Without asking price", x["unpriced"]),
            ]
        )
        out.append({"key": "inventory", "header": h, "rows": r, "notes": x["notes"]})
        dims = []
        for name, key in (
            ("Condition", "by_condition"),
            ("PTA", "by_pta"),
            ("ID check", "by_verification"),
            ("Type", "by_category"),
        ):
            dims += [
                [name, d["key"], d["count"], d["cost"], d["asking"]] for d in x[key] if d["count"]
            ]
        dims += [["Brand", b["brand"], b["count"], b["cost"], b["asking"]] for b in x["by_brand"]]
        out.append(
            {
                "key": "inventory_dims",
                "title": "Stock by dimension",
                "header": ["Dimension", "Value", "Phones", "At cost", "At asking"],
                "rows": dims,
                "notes": [],
            }
        )
    if "ageing" in want:
        x = inventory.ageing(db, shop_id, shop, can_cost=can_cost)
        out.append(
            {
                "key": "ageing",
                "header": ["Band", "Phones", "At cost", "At asking"],
                "rows": [[f"{b['band']}", b["count"], b["cost"], b["asking"]] for b in x["bands"]],
                "notes": x["notes"] + [f"Bands end at {x['band_edges']} days."],
            }
        )
        if detail:
            out.append(
                {
                    "key": "ageing_rows",
                    "title": "Every phone in stock",
                    "header": [
                        "Phone",
                        "IMEI",
                        "Supplier",
                        "In since",
                        "Days",
                        "Band",
                        "Cost",
                        "Asking",
                        "Margin %",
                    ],
                    "rows": [
                        [
                            r["label"],
                            r["imei1"],
                            r["supplier"],
                            r["intake_date"],
                            r["days"],
                            r["band"],
                            r["cost"],
                            r["asking"],
                            r["margin_pct"],
                        ]
                        for r in x["rows"]
                    ],
                    "notes": [],
                }
            )
        if x["alerts"]:
            out.append(
                {
                    "key": "alerts",
                    "title": "Needs attention",
                    "header": ["Alert", "Phone", "IMEI", "Days", "Cost", "Asking", "Supplier"],
                    "rows": [
                        [
                            a["kind"],
                            a["label"],
                            a["imei1"],
                            a["days"],
                            a["cost"],
                            a["asking"],
                            a["supplier"],
                        ]
                        for a in x["alerts"][:limit]
                    ],
                    "notes": [],
                }
            )
    if "products" in want:
        x = products.products(db, shop_id, shop, p, can_cost=can_cost, group="model")
        out.append(
            {
                "key": "products",
                "header": [
                    "Product",
                    "Units",
                    "Share %",
                    "Revenue",
                    "Cost",
                    "Profit",
                    "Margin %",
                    "Avg price",
                    "Avg days",
                    "Returns",
                    "In stock",
                    "Class",
                ],
                "rows": [
                    [
                        r["label"],
                        r["units"],
                        r["share"],
                        r["revenue"],
                        r["cost"],
                        r["profit"],
                        r["margin_pct"],
                        r["avg_price"],
                        r["avg_holding_days"],
                        r["returns"],
                        r["in_stock"],
                        r["performance"],
                    ]
                    for r in x["rows"][:limit]
                ],
                "notes": x["notes"],
            }
        )
    if "invoices" in want:
        x = products.invoices(db, shop_id, p, can_cost=can_cost)
        out.append(
            {
                "key": "invoices",
                "header": ["Kind", "Count", "Amount", "Voided"],
                "rows": [[k["kind"], k["count"], k["amount"], k["voided"]] for k in x["by_kind"]],
                "notes": x["notes"],
            }
        )
        out.append(
            {
                "key": "invoices_users",
                "title": "Documents by employee",
                "header": ["Employee", "Sales", "Sales amount", "Purchases", "Payments", "Voids"],
                "rows": [
                    [
                        u["user_name"],
                        u["sales"],
                        u["sales_amount"],
                        u["purchases"],
                        u["payments"],
                        u["voids"],
                    ]
                    for u in x["by_user"]
                ],
                "notes": [],
            }
        )
        if x["voided"]:
            out.append(
                {
                    "key": "voided",
                    "title": "Voided documents",
                    "header": ["Date", "Kind", "No", "Amount", "By", "Reason"],
                    "rows": [
                        [v["date"], v["kind"], v["no"], v["amount"], v["by"], v["reason"]]
                        for v in x["voided"][:limit]
                    ],
                    "notes": [],
                }
            )
    if "verification" in want:
        x = inventory.verification(db, shop_id, shop, p)
        h, r = _kv(
            [(f"In stock · {k}", v) for k, v in x["in_stock"].items()]
            + [(f"Checks saved · {k}", v) for k, v in x["checks_in_period"].items()]
            + [("Skipped checks", x["bypasses"])]
        )
        out.append({"key": "verification", "header": h, "rows": r, "notes": x["notes"]})
        if x["rows"]:
            out.append(
                {
                    "key": "bypasses",
                    "title": "Skipped ID checks",
                    "header": ["When", "Phone", "IMEI", "Supplier", "By", "Reason", "Was"],
                    "rows": [
                        [
                            b["at"],
                            b["label"],
                            b["imei1"],
                            b["party_name"],
                            b["user_name"],
                            b["reason"],
                            b["previous"],
                        ]
                        for b in x["rows"][:limit]
                    ],
                    "notes": [],
                }
            )
    if "checks" in want:
        rows = reconcile.checks(db, shop_id, p)
        out.append(
            {
                "key": "checks",
                "header": ["Check", "Left", "Right", "OK", "Difference"],
                "rows": [
                    [
                        c["key"] + (f" · {c['name']}" if c.get("name") else ""),
                        c["left"],
                        c["right"],
                        c["ok"],
                        c["diff"],
                    ]
                    for c in rows
                ],
                "notes": [],
            }
        )
    if "security" in want and is_owner:
        x = security.feed(db, shop_id, shop, p, limit=limit)
        out.append(
            {
                "key": "security",
                "header": [
                    "When",
                    "Kind",
                    "By",
                    "What",
                    "Before",
                    "After",
                    "Amount",
                    "Reason",
                    "Approved by",
                ],
                "rows": [
                    [
                        r["at"],
                        r["kind"],
                        r["user_name"],
                        r["label"],
                        r["before"],
                        r["after"],
                        r["amount"],
                        r["reason"],
                        r["approved_by"],
                    ]
                    for r in x["rows"]
                ],
                "notes": x["notes"],
            }
        )
        out.append(
            {
                "key": "security_users",
                "title": "Security events by employee",
                "header": ["Employee", "Events", "Flags", "Kinds"],
                "rows": [
                    [
                        u["user_name"],
                        u["total"],
                        ", ".join(u["flags"]),
                        ", ".join(f"{k} {v}" for k, v in u["kinds"].items()),
                    ]
                    for u in x["by_user"]
                ],
                "notes": [],
            }
        )
    for sec in out:
        sec.setdefault("title", TITLES.get(sec["key"], sec["key"]))
    return out


def xlsx(shop_name: str, p: Period, sections: list[dict], mode: str) -> Response:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Cover"
    ws.append(["Shop audit report"])
    ws.append(["Shop", shop_name])
    ws.append(["Period", f"{p.date_from} to {p.date_to}"])
    ws.append(["Balances as of", str(p.as_of)])
    ws.append(["Mode", mode])
    ws.append(["Currency", "PKR"])
    ws.append(["Generated", dt.datetime.now().strftime("%d-%m-%Y %H:%M")])
    ws.append(["Sections", ", ".join(s["title"] for s in sections)])
    ws["A1"].font = Font(bold=True, size=14)
    used: set[str] = set()
    for sec in sections:
        name = sec["title"][:28]
        base, i = name, 2
        while name in used:
            name, i = f"{base[:25]} {i}", i + 1
        used.add(name)
        w = wb.create_sheet(title="".join(ch for ch in name if ch not in "[]:*?/\\"))
        w.append(sec["header"])
        for cell in w[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="F6F4EF")
        for r in sec["rows"]:
            w.append(
                [
                    float(v)
                    if isinstance(v, Decimal)
                    else (_s(v) if isinstance(v, dt.date | dt.datetime | bool) else v)
                    for v in r
                ]
            )
        if sec["notes"]:
            w.append([])
            for note in sec["notes"]:
                w.append([note])
    buf = io.BytesIO()
    wb.save(buf)
    stamp = dt.date.today().isoformat()
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="shop-audit-{stamp}.xlsx"'},
    )


def printable(
    shop_name: str, p: Period, sections: list[dict], mode: str, lang: str | None
) -> Response:
    e = html.escape
    ur = lang == "ur"
    parts = []
    for sec in sections:
        head = "".join(f"<th>{e(str(h))}</th>" for h in sec["header"])
        body = (
            "".join(
                "<tr>"
                + "".join(
                    f'<td class="{"n" if isinstance(v, int | float | Decimal) else ""}" dir="auto">{e(_s(v))}</td>'
                    for v in r
                )
                + "</tr>"
                for r in sec["rows"]
            )
            or '<tr><td colspan="99" class="muted">—</td></tr>'
        )
        notes = "".join(f"<li>{e(n_)}</li>" for n_ in sec["notes"])
        parts.append(
            f"<section><h2>{e(sec['title'])}</h2><table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>{f'<ul class=notes>{notes}</ul>' if notes else ''}</section>"
        )
    now = dt.datetime.now().strftime("%d-%m-%Y %I:%M %p")
    page = f"""<!doctype html><html lang="{"ur" if ur else "en"}" dir="{"rtl" if ur else "ltr"}"><head><meta charset="utf-8"><title>Shop audit · {e(shop_name)}</title>
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Manrope:wght@500;700;800&family=Noto+Nastaliq+Urdu&display=swap">
<style>
@page {{ size: A4 portrait; margin: 12mm; }}
body {{ font-family: Manrope, 'Noto Nastaliq Urdu', sans-serif; color: #1B1F24; margin: 0; }}
h1 {{ font-size: 22px; margin: 0; }} h2 {{ font-size: 15px; margin: 18px 0 6px; page-break-after: avoid; }}
.meta {{ color: #5E646B; font-size: 12px; margin: 4px 0 14px; }}
table {{ width: 100%; border-collapse: collapse; font-size: 11px; page-break-inside: auto; }}
th {{ background: #F6F4EF; text-align: start; }} th, td {{ border: 1px solid #E6E2DA; padding: 5px 7px; }}
td.n {{ text-align: right; font-variant-numeric: tabular-nums; white-space: nowrap; }}
tr:nth-child(even) td {{ background: #FBFAF7; }} tr {{ page-break-inside: avoid; }}
.notes {{ color: #5E646B; font-size: 10px; margin: 6px 0 0; padding-inline-start: 16px; }} .muted {{ color: #5E646B; }}
section {{ page-break-inside: avoid; }}
</style></head><body>
<h1>{e(shop_name)} · Shop audit report</h1>
<div class="meta">Period {e(str(p.date_from))} → {e(str(p.date_to))} · balances as of {e(str(p.as_of))} · {e(mode)} · PKR · generated {e(now)} · {len(sections)} sections</div>
{"".join(parts)}
</body></html>"""
    return Response(content=page, media_type="text/html; charset=utf-8")
