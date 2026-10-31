"""Read an uploaded .xlsx or .csv into rows keyed by column key. Headers are matched tolerantly
(English or Urdu header, or the key itself; case, spaces and the * mark do not matter)."""

import csv
import datetime as dt
import io
from decimal import Decimal, InvalidOperation

from app.core.errors import ValidationFailed
from app.core.pk import latin_digits
from app.modules.bulk.spec import COLUMNS, Col, norm

MAX_ROWS = 2000
MAX_BYTES = 5 * 1024 * 1024
NOTE_MARKS = ("↑", "←")  # ↑ ←
SAMPLE_FILL = "FFF4B3"  # the template paints its sample rows this yellow
SAMPLE_MARK = object()


def _header_map(kind: str, headers: list[str]) -> dict[int, str]:
    cols: list[Col] = COLUMNS[kind]
    wanted = {}
    for c in cols:
        for h in (c.key, c.en, c.ur):
            wanted[norm(h.replace("*", ""))] = c.key
    out: dict[int, str] = {}
    for i, h in enumerate(headers):
        k = wanted.get(norm(str(h or "").replace("*", "")))
        if k and k not in out.values():
            out[i] = k
    return out


def read_rows(kind: str, data: bytes, filename: str | None) -> list[dict]:
    if len(data) > MAX_BYTES:
        raise ValidationFailed(
            code="file_too_big",
            message_en="File is larger than 5 MB",
            message_ur="فائل 5 MB سے بڑی ہے",
        )
    name = (filename or "").lower()
    if name.endswith(".csv") or (not name.endswith(".xlsx") and data[:2] != b"PK"):
        text = data.decode("utf-8-sig", errors="replace")
        table = [row for row in csv.reader(io.StringIO(text))]
    else:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), data_only=True)
        ws = wb["Data"] if "Data" in wb.sheetnames else wb.worksheets[0]
        table = []
        for row in ws.iter_rows(max_row=min(ws.max_row, MAX_ROWS + 12)):
            vals = [c.value for c in row]
            first = row[0] if row else None
            rgb = getattr(getattr(getattr(first, "fill", None), "fgColor", None), "rgb", None)
            if isinstance(rgb, str) and rgb.upper().endswith(SAMPLE_FILL):
                vals.append(SAMPLE_MARK)  # a yellow row: one of the template's samples
            table.append(vals)
    while table and all(v in (None, "") or v is SAMPLE_MARK for v in table[-1]):
        table.pop()  # trailing empty lines (Excel keeps formatted blanks)
    # the header is the first row that matches at least two of our columns
    header_idx = None
    hmap: dict[int, str] = {}
    for i, row in enumerate(table[:10]):
        m = _header_map(kind, [str(c) if c is not None else "" for c in row])
        if len(m) >= 2:
            header_idx, hmap = i, m
            break
    if header_idx is None:
        raise ValidationFailed(
            code="no_header",
            message_en="Could not find the column headers. Use the downloaded template.",
            message_ur="کالم کے عنوان نہیں ملے۔ ڈاؤن لوڈ کیا ہوا ٹیمپلیٹ استعمال کریں۔",
        )
    rows: list[dict] = []
    for n, row in enumerate(table[header_idx + 1 :], start=header_idx + 2):
        if _is_note(row):
            continue
        rec = {k: row[i] if i < len(row) else None for i, k in hmap.items()}
        rows.append({"_row": n, "_sample": bool(row) and row[-1] is SAMPLE_MARK, **rec})
    if len(rows) > MAX_ROWS:
        raise ValidationFailed(
            code="too_many_rows",
            message_en=f"At most {MAX_ROWS} rows per file",
            message_ur=f"ایک فائل میں زیادہ سے زیادہ {MAX_ROWS} قطاریں",
        )
    return rows


def _is_note(row) -> bool:
    """the template's own hint line under the samples: an arrow in the first cell, nothing else"""
    first = str(row[0] or "").strip() if row else ""
    rest = [c for c in row[1:] if c not in (None, "") and c is not SAMPLE_MARK]
    return bool(first) and first[0] in NOTE_MARKS and not rest


def cell_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, dt.datetime | dt.date):
        return v.strftime("%d-%m-%Y")
    return str(v).strip()


def cell_money(v) -> Decimal | None:
    """'35,000', '35000.50', 'Rs. 35000', Urdu digits → Decimal; '' → None; junk → raises"""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, int | float | Decimal):
        return Decimal(str(v))
    s = (
        latin_digits(str(v))
        .replace(",", "")
        .replace("Rs.", "")
        .replace("Rs", "")
        .replace("روپے", "")
        .strip()
    )
    try:
        return Decimal(s)
    except InvalidOperation as e:
        raise ValueError("number") from e


def cell_int(v) -> int | None:
    d = cell_money(v)
    if d is None:
        return None
    if d != d.to_integral_value():
        raise ValueError("int")
    return int(d)


def cell_date(v) -> dt.date | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = latin_digits(str(v)).strip().replace("/", "-").replace(".", "-")
    for fmt in ("%d-%m-%Y", "%Y-%m-%d", "%d-%m-%y", "%d %b %Y", "%d %B %Y"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError("date")
