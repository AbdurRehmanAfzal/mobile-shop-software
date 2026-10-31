"""The workbook a shopkeeper downloads: Data (headers + three yellow sample rows + dropdowns),
Help (every column explained), Lists (what the dropdowns read). Also the round-trip export in
the same layout."""

import datetime as dt
import io
import uuid

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.modules.bulk.spec import COLUMNS, ENUM_LABELS, SAMPLES, Col
from app.modules.catalogue.models import Brand, Colour, Model, Storage
from app.modules.parties.models import Party

YELLOW = PatternFill("solid", fgColor="FFF4B3")
HEAD = PatternFill("solid", fgColor="E6E2DA")


def _visible(col, shop_id):
    return (col.is_(None)) | (col == shop_id)


def lists_for(db: Session, shop_id: uuid.UUID, kind: str, lang: str | None) -> dict[str, list[str]]:
    ur = lang == "ur"
    out: dict[str, list[str]] = {}
    if kind == "stock":
        brands = (
            db.execute(
                select(Brand.name)
                .where(_visible(Brand.shop_id, shop_id), Brand.is_active.is_(True))
                .order_by(Brand.name)
            )
            .scalars()
            .all()
        )
        out["brand"] = list(brands)
        models = db.execute(
            select(Brand.name, Model.name)
            .join(Brand, Brand.id == Model.brand_id)
            .where(_visible(Model.shop_id, shop_id), Model.is_active.is_(True))
            .order_by(Brand.name, Model.name)
        ).all()
        out["model"] = [f"{b} · {m}" for b, m in models]
        out["storage"] = list(
            db.execute(
                select(Storage.name)
                .where(_visible(Storage.shop_id, shop_id))
                .order_by(Storage.gigabytes)
            )
            .scalars()
            .all()
        )
        out["colour"] = list(
            db.execute(
                select(Colour.name).where(_visible(Colour.shop_id, shop_id)).order_by(Colour.name)
            )
            .scalars()
            .all()
        )
        out["supplier"] = [
            f"{p.name} · {p.phone}"
            for p in db.scalars(
                select(Party)
                .where(Party.shop_id == shop_id, Party.is_active.is_(True))
                .order_by(Party.name)
            ).all()
            if "supplier" in (p.kinds or [])
        ]
    if kind == "items":
        from app.modules.accessories.models import CATEGORIES

        out["item_category"] = list(CATEGORIES)
    if kind == "expenses":
        from app.modules.expenses.service import list_categories
        from app.modules.ledger.service import list_cash_accounts

        out["expense_category"] = [
            (c["name_ur"] if ur and c.get("name_ur") else c["name"])
            for c in list_categories(db, shop_id)
        ]
        out["account"] = [a["name"] for a in list_cash_accounts(db, shop_id)]
    for c in COLUMNS[kind]:
        if c.kind == "enum" and c.enum in ENUM_LABELS:
            out[c.enum] = [lbl[1] if ur else lbl[0] for lbl in ENUM_LABELS[c.enum].values()]
    return out


def build(
    db: Session, shop_id: uuid.UUID, kind: str, lang: str | None, *, rows: list[list] | None = None
) -> bytes:
    """`rows=None` → the template with sample rows; `rows=[...]` → an export in the same layout."""
    ur = lang == "ur"
    cols: list[Col] = COLUMNS[kind]
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.sheet_view.rightToLeft = ur
    headers = [(c.ur if ur else c.en) + (" *" if c.required else "") for c in cols]
    ws.append(headers)
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True)
        cell.fill = HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(28, len(c.en) + 6))
    ws.freeze_panes = "A2"
    if rows is None:
        for s in SAMPLES[kind]:
            ws.append([_sample_cell(s.get(c.key), c, ur) for c in cols])
        for r in range(2, 2 + len(SAMPLES[kind])):
            for i in range(1, len(cols) + 1):
                ws.cell(row=r, column=i).fill = YELLOW
        note = ws.cell(row=2 + len(SAMPLES[kind]), column=1)
        note.value = (
            "← نمونہ قطاریں (پیلی): اپنا ڈیٹا ان کی جگہ یا ان کے نیچے لکھیں؛ یہ خود چھوڑ دی جائیں گی"
            if ur
            else "↑ Sample rows (yellow): write your data in their place or below; they are skipped automatically"
        )
        note.font = Font(italic=True, color="5E646B")
    else:
        for r in rows:
            ws.append(r)
    # dropdowns for the enum / list columns
    lists = lists_for(db, shop_id, kind, lang)
    wl = wb.create_sheet("Lists")
    wl.sheet_view.rightToLeft = ur
    col_i = 1
    for c in cols:
        key = c.enum if c.kind == "enum" else c.key if c.kind == "list" else None
        values = lists.get(key or "", [])
        if not key or not values:
            continue
        wl.cell(row=1, column=col_i, value=(c.ur if ur else c.en)).font = Font(bold=True)
        for j, v in enumerate(values[:500], start=2):
            wl.cell(row=j, column=col_i, value=v)
        letter = get_column_letter(col_i)
        wl.column_dimensions[letter].width = max(
            14, min(36, max(len(str(v)) for v in values[:500]) + 2)
        )
        if c.kind == "enum" or c.key in ("brand", "storage", "colour", "account"):
            dv = DataValidation(
                type="list",
                formula1=f"=Lists!${letter}$2:${letter}${len(values[:500]) + 1}",
                allow_blank=True,
                showErrorMessage=False,
            )
            ws.add_data_validation(dv)
            idx = cols.index(c) + 1
            dv.add(f"{get_column_letter(idx)}2:{get_column_letter(idx)}2001")
        col_i += 1
    # help
    wh = wb.create_sheet("Help")
    wh.sheet_view.rightToLeft = ur
    wh.append(["کالم", "ضروری؟", "کیا لکھیں"] if ur else ["Column", "Required?", "What to write"])
    for i in range(1, 4):
        wh.cell(row=1, column=i).font = Font(bold=True)
    for c in cols:
        wh.append(
            [
                (c.ur if ur else c.en),
                ("ہاں" if ur else "yes") if c.required else ("نہیں" if ur else "no"),
                (c.help_ur if ur else c.help_en) or "",
            ]
        )
    wh.append([])
    wh.append(
        [
            (
                "تاریخیں دن-مہینہ-سال؛ رقم سادہ ہندسوں میں؛ ایک قطار = ایک "
                + ("موبائل" if kind in ("stock", "sales") else "اندراج")
            )
            if ur
            else (
                "Dates dd-mm-yyyy; money in plain rupees; one row = one "
                + ("phone" if kind in ("stock", "sales") else "entry")
            )
        ]
    )
    wh.append(
        [
            (
                "اپ لوڈ کے بعد پہلے جائزہ آتا ہے — کچھ بھی محفوظ نہیں ہوتا جب تک آپ 'درآمد کریں' نہ دبائیں۔"
            )
            if ur
            else ("After upload you see a preview first — nothing is saved until you press Import.")
        ]
    )
    wh.column_dimensions["A"].width = 26
    wh.column_dimensions["B"].width = 10
    wh.column_dimensions["C"].width = 70
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _sample_cell(v, c: Col, ur: bool):
    if v is None:
        return None
    if c.kind == "enum" and c.enum in ENUM_LABELS:
        from app.modules.bulk.spec import enum_value

        key = enum_value(c.enum, v)
        if key:
            return ENUM_LABELS[c.enum][key][1 if ur else 0]
    return v


def filename(kind: str, export: bool = False) -> str:
    stamp = dt.date.today().isoformat()
    return f"{kind}-{'export' if export else 'template'}-{stamp}.xlsx"
