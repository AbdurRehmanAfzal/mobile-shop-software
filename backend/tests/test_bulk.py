"""Owner 2026-10-31 (docs/18): bulk Excel templates with sample rows, a preview that writes nothing,
an all-or-nothing import through the normal services, and an export in the same layout."""

import io
from decimal import Decimal

from openpyxl import load_workbook

from tests.conftest import login
from tests.test_purchases import _dev, _party

A, B, C = "490154203237518", "356938035643809", "352099001761481"  # the template's sample IMEIs
D, E, F = "359568073814231", "861536030196001", "353915057311922"  # fresh ones (Luhn-valid)


def _wb(data: bytes):
    return load_workbook(io.BytesIO(data))


def _csv(text: str, name: str = "rows.csv") -> dict:
    return {"files": {"file": (name, text.encode("utf-8"), "text/csv")}}


def _xlsx(headers: list[str], rows: list[list], name: str = "rows.xlsx") -> dict:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(headers)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return {"files": {"file": (name, buf.getvalue(), "application/octet-stream")}}


def _imeis(actor) -> set[str]:
    return {d["imei1"] for d in actor.get("/api/v1/devices").json()}


# ---- templates --------------------------------------------------------------------------------


def test_templates_have_data_help_lists_and_yellow_sample_rows(owner, cat):
    kinds = owner.get("/api/v1/bulk/kinds").json()
    assert [k["kind"] for k in kinds] == ["stock", "people", "items", "sales", "expenses"]
    for kind in ("stock", "people", "items", "sales", "expenses"):
        r = owner.get(f"/api/v1/bulk/template/{kind}")
        assert r.status_code == 200, r.text
        assert r.headers["content-disposition"].startswith(
            f'attachment; filename="{kind}-template-'
        )
        wb = _wb(r.content)
        assert wb.sheetnames == ["Data", "Lists", "Help"]
        ws = wb["Data"]
        assert ws.max_row >= 5  # header + 3 samples + the note
        assert ws["A1"].font.bold
        assert ws["A2"].fill.fgColor.rgb.endswith("FFF4B3") and ws["A4"].fill.fgColor.rgb.endswith(
            "FFF4B3"
        )
        assert "Sample rows" in str(ws.cell(row=5, column=1).value)
        assert wb["Help"].max_row > 4
    # stock: the sheet carries the required marks, dropdown lists and seeded brands
    wb = _wb(owner.get("/api/v1/bulk/template/stock").content)
    heads = [c.value for c in wb["Data"][1]]
    assert heads[:4] == ["IMEI 1 *", "IMEI 2", "Brand *", "Model *"]
    assert wb["Data"]["A2"].value == A and wb["Data"]["G2"].value == "Used"
    lists = [c.value for c in wb["Lists"][1]]
    assert {"Brand", "Condition", "PTA"} <= set(lists)
    brand_col = lists.index("Brand") + 1
    names = {wb["Lists"].cell(row=i, column=brand_col).value for i in range(2, 40)}
    assert "Nokia" in names
    assert wb["Data"].data_validations.dataValidation  # dropdowns are attached
    # Urdu template: Urdu headers, right-to-left sheet, Urdu enum words in the samples
    wb = _wb(owner.get("/api/v1/bulk/template/stock", params={"lang": "ur"}).content)
    assert wb["Data"]["C1"].value == "برانڈ *" and wb["Data"].sheet_view.rightToLeft
    assert wb["Data"]["G2"].value == "استعمال شدہ"


def test_template_uploaded_unchanged_only_skips_the_samples(owner, cat):
    data = owner.get("/api/v1/bulk/template/stock").content
    r = owner.post(
        "/api/v1/bulk/preview/stock",
        files={"file": ("stock-template.xlsx", data, "application/octet-stream")},
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert (
        body["counts"]["skip"] >= 3 and body["counts"]["ok"] == 0 and body["counts"]["error"] == 0
    )
    assert body["can_import"] is False
    assert all(x["status"] == "skip" for x in body["rows"])
    assert "Sample row" in body["rows"][0]["messages"][0]
    assert _imeis(owner) == set()


# ---- stock ------------------------------------------------------------------------------------


STOCK_CSV = (
    "IMEI 1,Brand,Model,Storage,Colour,Condition,PTA,SIM time (months),Cost (Rs.),Asking price (Rs.),"
    "Supplier phone,Supplier name,Date,Paid now (Rs.),Notes\n"
    f"{D},Nokia,105 4G,,,New,PTA approved,,3000,3500,,,,,own stock\n"
    f"{E},Zyrex,Z1,128 GB,Black,Used,Non-PTA,4,22000,26000,0300-7778899,Rehman Traders,15-10-2026,20000,\n"
)


def test_stock_preview_writes_nothing_and_import_posts_everything(owner, cat):
    pv = owner.post("/api/v1/bulk/preview/stock", **_csv(STOCK_CSV)).json()
    assert pv["counts"] == {"ok": 1, "warn": 1, "error": 0, "skip": 0}
    assert pv["can_import"] is True and pv["already_imported_at"] is None
    r1, r2 = pv["rows"]
    assert r1["action"] == "opening_stock" and r1["values"]["condition"] == "New"
    assert r2["action"] == "purchase" and r2["values"]["supplier_phone"] == "03007778899"
    assert any("Brand 'Zyrex' is new" in m for m in r2["messages"])
    assert any("New person 'Rehman Traders'" in m for m in r2["messages"])
    assert _imeis(owner) == set() and owner.get("/api/v1/parties").json() == []
    assert owner.get("/api/v1/bulk/imports").json() == []

    r = owner.post("/api/v1/bulk/import/stock", **_csv(STOCK_CSV))
    assert r.status_code == 201, r.text
    body = r.json()
    assert body["created"] == {"devices": 2, "brands": 1, "models": 1, "parties": 1, "purchases": 1}
    assert _imeis(owner) == {D, E}
    devs = {d["imei1"]: d for d in owner.get("/api/v1/devices").json()}
    assert devs[D]["cost_price"] == "3000.00" and devs[D]["asking_price"] == "3500.00"
    assert devs[E]["brand_name"] == "Zyrex" and devs[E]["pta_status"] == "non_pta"
    assert devs[E]["sim_time_months"] == 4 and devs[E]["storage_name"] == "128GB"
    pur = owner.get("/api/v1/purchases").json()
    pur = pur["items"] if isinstance(pur, dict) else pur
    assert len(pur) == 1 and pur[0]["paid"] == "20000.00" and pur[0]["date"] == "2026-10-15"
    sup = owner.get("/api/v1/parties").json()[0]
    assert sup["phone"] == "03007778899" and "supplier" in sup["kinds"]
    assert Decimal(sup["balance"]) == Decimal("-2000.00")  # we still owe 2,000
    log = owner.get("/api/v1/bulk/imports").json()
    assert log[0]["kind"] == "stock" and log[0]["rows_ok"] == 2 and log[0]["by"] == "Owner"

    # the same file again: every IMEI is now taken, and the file is remembered
    pv = owner.post("/api/v1/bulk/preview/stock", **_csv(STOCK_CSV)).json()
    assert pv["counts"]["error"] == 2 and pv["already_imported_at"] is not None
    assert "already in your shop" in pv["rows"][0]["messages"][0]


def test_stock_import_is_all_or_nothing_unless_bad_rows_are_skipped(owner, cat):
    bad = STOCK_CSV + "12345,Nokia,105,,,New,PTA approved,,100,,,,,,\n"
    r = owner.post("/api/v1/bulk/import/stock", **_csv(bad))
    assert r.status_code == 422 and r.json()["error"]["code"] == "bulk_has_errors"
    details = r.json()["error"]["details"]
    assert details["counts"]["error"] == 1 and "15 digits" in details["rows"][2]["messages"][0]
    assert _imeis(owner) == set() and owner.get("/api/v1/parties").json() == []

    r = owner.post("/api/v1/bulk/import/stock", data={"skip_errors": "true"}, **_csv(bad))
    assert r.status_code == 201, r.text
    assert r.json()["counts"] == {"ok": 1, "warn": 1, "error": 1, "skip": 0}
    assert _imeis(owner) == {D, E}


def test_stock_rows_are_checked_in_the_shopkeepers_words(owner, cat):
    csv = (
        "IMEI 1,Brand,Model,Condition,PTA,Cost (Rs.),Asking price (Rs.),Supplier phone,Supplier name\n"
        f"{D},Nokia,105,Brand new,PTA,abc,,,\n"  # cost not a number
        f"{D},Nokia,105,Used,Maybe,100,90,,\n"  # duplicate IMEI, bad PTA word, asking below cost
        f"{E},Nokia,105,Used,JV,100,,0301,\n"  # bad phone
        f",Nokia,105,Used,JV,100,,,\n"  # missing IMEI
        f"{C},Nokia,105,New,PTA approved,100,,,\n"  # sample IMEI → skipped
        ",,,,,,,,\n"
    )
    pv = owner.post("/api/v1/bulk/preview/stock", **_csv(csv)).json()
    rows = {r["row"]: r for r in pv["rows"]}
    assert "Cost (Rs.) must be a number" in rows[2]["messages"]
    assert rows[2]["values"]["condition"] == "New"  # "Brand new" understood
    assert "This IMEI appears twice in the file" in rows[3]["messages"]
    assert any(m.startswith("PTA: write one of") for m in rows[3]["messages"])
    assert "Asking price is below cost" in rows[3]["messages"]
    assert any("Pakistani mobile number" in m for m in rows[4]["messages"])
    assert "IMEI 1 is required" in rows[5]["messages"]
    assert rows[6]["status"] == "skip" and 7 not in rows  # the trailing blank line is dropped
    # the same file in Urdu
    pv = owner.post("/api/v1/bulk/preview/stock", params={"lang": "ur"}, **_csv(csv)).json()
    assert "ضروری ہے" in pv["rows"][3]["messages"][0]
    # without "create new brands" an unknown brand is an error
    pv = owner.post(
        "/api/v1/bulk/preview/stock", data={"create_catalogue": "false"}, **_csv(STOCK_CSV)
    ).json()
    assert pv["rows"][1]["status"] == "error" and "not in your list" in pv["rows"][1]["messages"][0]


def test_xlsx_upload_with_urdu_headers_and_excel_dates(owner, cat):
    import datetime as dt

    up = _xlsx(
        ["IMEI 1", "برانڈ", "ماڈل", "حالت", "PTA", "لاگت (روپے)", "تاریخ"],
        [[int(D), "Nokia", "105 4G", "استعمال شدہ", "PTA منظور", 3000.0, dt.datetime(2026, 10, 2)]],
    )
    r = owner.post("/api/v1/bulk/import/stock", **up)
    assert r.status_code == 201, r.text
    assert r.json()["rows"][0]["values"]["date"] == "02-10-2026"
    assert _imeis(owner) == {D}


# ---- people -----------------------------------------------------------------------------------


def test_people_import_creates_parties_with_opening_balances_and_skips_known_phones(owner):
    _party(owner, name="Known", phone="03009990000", kinds=("customer",))
    csv = (
        "Name,Phone,Type,City,CNIC,Opening balance (Rs.),Who owes,Notes\n"
        "Ahmed Khan,0300 1112233,Customer,Lahore,33100-1234567-1,15000,They owe us,regular\n"
        "Rehman Traders,+92 321 1112233,Supplier,Karachi,,60000,We owe them,\n"
        "Bilal Mobile,03451112233,Customer & supplier,Faisalabad,,,,\n"
        "Known Again,03009990000,Customer,,,,,\n"
        "Ahmed Khan,03001234567,Customer,,,,,\n"  # a sample phone → skipped
        "Lost,03121112233,Customer,,,500,,\n"  # opening without a direction
    )
    pv = owner.post("/api/v1/bulk/preview/people", **_csv(csv)).json()
    assert pv["counts"] == {"ok": 3, "warn": 0, "error": 1, "skip": 2}
    assert "Uses the existing person 'Known'" in pv["rows"][3]["messages"][0]
    assert "Say who owes" in pv["rows"][5]["messages"][0]
    r = owner.post("/api/v1/bulk/import/people", data={"skip_errors": "true"}, **_csv(csv))
    assert r.status_code == 201, r.text
    assert r.json()["created"] == {"parties": 3, "opening_balances": 2}
    by = {p["phone"]: p for p in owner.get("/api/v1/parties").json()}
    assert (
        by["03001112233"]["balance"] == "15000.00"
        and by["03001112233"]["cnic"] == "33100-1234567-1"
    )
    assert by["03211112233"]["balance"] == "-60000.00" and by["03211112233"]["kinds"] == [
        "supplier"
    ]
    assert set(by["03451112233"]["kinds"]) == {"customer", "supplier"}
    assert by["03009990000"]["name"] == "Known"
    assert "03121112233" not in by


# ---- items ------------------------------------------------------------------------------------


def test_items_import(owner):
    csv = (
        "Name,Category,SKU,Quantity,Unit cost (Rs.),Sale price (Rs.),Reorder at\n"
        "Fast charger 33W,Charger,,20,450,800,5\n"
        "Tempered glass (Y21),glass,TG-Y21,50,60,250,\n"
        "Mystery,gadget,TG-Y21,1,1,2,\n"  # unknown category (warning) + duplicate SKU (error)
    )
    pv = owner.post("/api/v1/bulk/preview/items", **_csv(csv)).json()
    assert pv["counts"] == {"ok": 2, "warn": 0, "error": 1, "skip": 0}
    assert pv["rows"][0]["values"]["category"] == "charger"
    assert any("Category 'gadget'" in m for m in pv["rows"][2]["messages"])
    assert any("SKU 'TG-Y21' already exists" in m for m in pv["rows"][2]["messages"])
    r = owner.post("/api/v1/bulk/import/items", data={"skip_errors": "true"}, **_csv(csv))
    assert r.status_code == 201, r.text
    items = {a["name"]: a for a in owner.get("/api/v1/accessories").json()}
    assert (
        items["Fast charger 33W"]["qty"] == 20 and items["Fast charger 33W"]["avg_cost"] == "450.00"
    )
    assert (
        items["Tempered glass (Y21)"]["sku"] == "TG-Y21"
        and items["Tempered glass (Y21)"]["reorder_level"] == 5
    )


# ---- sales ------------------------------------------------------------------------------------


def test_sales_import_sells_stock_phones_with_khata_and_below_cost_guard(owner, cat):
    sup = _party(owner, name="Dealer", phone="03005556666", kinds=("supplier",))
    owner.post(
        "/api/v1/purchases",
        json={
            "kind": "supplier",
            "party_id": sup["id"],
            "devices": [
                {"device": _dev(cat, D), "cost": 800, "asking_price": 1000},
                {"device": _dev(cat, E), "cost": 800, "asking_price": 1000},
            ],
            "paid": 1600,
        },
    )
    csv = (
        "IMEI,Sold for (Rs.),Paid now (Rs.),Customer phone,Customer name,Date,Notes\n"
        f"{D},1000,600,03001112233,Ahmed Khan,20-10-2026,rest on khata\n"
        f"{E},700,700,,,,\n"  # below cost
        f"{A},1000,1000,,,,\n"  # sample → skipped
        f"{F},1000,1000,,,,\n"  # not in stock
        f"{D},1000,900,,,,\n"  # duplicate in file; also khata without a customer
    )
    pv = owner.post("/api/v1/bulk/preview/sales", **_csv(csv)).json()
    rows = pv["rows"]
    assert rows[0]["status"] == "warn" and "New person 'Ahmed Khan'" in rows[0]["messages"][0]
    assert rows[1]["status"] == "error" and "Sold under cost (800)" in rows[1]["messages"][0]
    assert rows[2]["status"] == "skip"
    assert "No phone with this IMEI is in stock" in rows[3]["messages"]
    assert "This IMEI appears twice in the file" in rows[4]["messages"]
    assert "Something stays on khata: customer phone is needed" in rows[4]["messages"]
    assert _imeis(owner) == {D, E}
    # with the loss confirmed by the owner, the below-cost row passes
    pv = owner.post(
        "/api/v1/bulk/preview/sales", data={"allow_below_cost": "true"}, **_csv(csv)
    ).json()
    assert pv["rows"][1]["status"] == "warn"
    r = owner.post(
        "/api/v1/bulk/import/sales",
        data={"allow_below_cost": "true", "skip_errors": "true"},
        **_csv(csv),
    )
    assert r.status_code == 201, r.text
    assert r.json()["created"] == {"parties": 1, "invoices": 2}
    devs = {d["imei1"]: d for d in owner.get("/api/v1/devices", params={"status": "sold"}).json()}
    assert set(devs) == {D, E}
    cus = next(p for p in owner.get("/api/v1/parties").json() if p["phone"] == "03001112233")
    assert cus["balance"] == "400.00" and "customer" in cus["kinds"]
    inv = owner.get("/api/v1/sales").json()
    inv = inv["items"] if isinstance(inv, dict) else inv
    assert {i["date"] for i in inv} >= {"2026-10-20"}
    # the same file again: both phones are sold now
    pv = owner.post("/api/v1/bulk/preview/sales", **_csv(csv)).json()
    assert "Phone is sold, not in stock" in pv["rows"][0]["messages"]


def test_sales_rows_cannot_see_another_shops_phones(owner, other_owner, cat):
    sup = _party(owner, name="Dealer", phone="03005556666", kinds=("supplier",))
    owner.post(
        "/api/v1/purchases",
        json={
            "kind": "supplier",
            "party_id": sup["id"],
            "devices": [{"device": _dev(cat, D), "cost": 800}],
            "paid": 800,
        },
    )
    pv = other_owner.post(
        "/api/v1/bulk/preview/sales", **_csv(f"IMEI,Sold for (Rs.),Paid now (Rs.)\n{D},1000,1000\n")
    ).json()
    assert (
        pv["rows"][0]["status"] == "error"
        and "No phone with this IMEI" in pv["rows"][0]["messages"][0]
    )


# ---- expenses ---------------------------------------------------------------------------------


def test_expenses_import_matches_categories_and_accounts(owner):
    owner.post("/api/v1/ledger/accounts", json={"kind": "bank", "name": "HBL"})
    csv = (
        "Date,Category,Amount (Rs.),Description,Account\n"
        "01-10-2026,Shop rent,25500,October,\n"
        '05/10/2026,electricity,"6,200",,HBL\n'
        "10-10-2026,چائے پانی,800,,Cash\n"
        "11-10-2026,Donations,500,,\n"  # unknown category
        "32-10-2026,Shop rent,500,,\n"  # bad date
        "12-10-2026,Shop rent,0,,Meezan\n"  # zero amount + unknown account
    )
    pv = owner.post("/api/v1/bulk/preview/expenses", **_csv(csv)).json()
    assert pv["counts"] == {"ok": 3, "warn": 0, "error": 3, "skip": 0}
    assert (
        pv["rows"][1]["values"]["amount"] == "6200"
        and pv["rows"][1]["values"]["date"] == "05-10-2026"
    )
    assert pv["rows"][2]["values"]["category"] == "Tea & refreshments"
    assert "Category 'Donations' is not in your list" in pv["rows"][3]["messages"]
    assert "Date: write the date as dd-mm-yyyy" in pv["rows"][4]["messages"]
    assert any("Account 'Meezan' not found" in m for m in pv["rows"][5]["messages"])
    r = owner.post("/api/v1/bulk/import/expenses", data={"skip_errors": "true"}, **_csv(csv))
    assert r.status_code == 201, r.text
    assert r.json()["created"] == {"expenses": 3}
    exp = owner.get("/api/v1/expenses").json()
    exp = exp["items"] if isinstance(exp, dict) else exp
    assert len(exp) == 3
    accounts = {a["name"]: a for a in owner.get("/api/v1/ledger/accounts").json()}
    assert accounts["HBL"]["balance"] == "-6200.00"
    assert accounts["Cash"]["balance"] == "-26300.00"


# ---- export, roles, limits --------------------------------------------------------------------


def test_export_uses_the_template_layout_and_round_trips(owner, cat):
    owner.post("/api/v1/bulk/import/stock", **_csv(STOCK_CSV))
    r = owner.get("/api/v1/bulk/export/stock")
    assert r.status_code == 200 and r.headers["content-disposition"].startswith(
        'attachment; filename="stock-export-'
    )
    wb = _wb(r.content)
    ws = wb["Data"]
    assert [c.value for c in ws[1]][:3] == ["IMEI 1 *", "IMEI 2", "Brand *"]
    assert {ws.cell(row=i, column=1).value for i in (2, 3)} == {D, E}
    assert ws.cell(row=3, column=3).value == "Zyrex" and ws.cell(row=3, column=8).value == "Non-PTA"
    assert ws.cell(row=3, column=10).value == 22000
    # the export of another kind as csv
    r = owner.get("/api/v1/bulk/export/people", params={"format": "csv"})
    assert r.status_code == 200 and "Rehman Traders,03007778899,Supplier" in r.content.decode(
        "utf-8-sig"
    )
    r = owner.get("/api/v1/bulk/export/items")
    assert r.status_code == 200 and _wb(r.content)["Data"].max_row >= 1


def test_only_owner_and_manager_use_bulk(client, owner):
    owner.post(
        "/api/v1/staff",
        json={"name": "Sales", "phone": "03451112233", "password": "secret123", "role": "salesman"},
    )
    s = login(client, "03451112233")
    assert s.get("/api/v1/bulk/template/stock").status_code == 403
    assert (
        s.post(
            "/api/v1/bulk/preview/people", **_csv("Name,Phone,Type\nA B,03001112233,Customer\n")
        ).status_code
        == 403
    )
    assert owner.get("/api/v1/bulk/template/nothing").status_code == 404


def test_bad_files_are_refused_kindly(owner):
    r = owner.post("/api/v1/bulk/preview/people", **_csv("hello,world\n1,2\n"))
    assert r.status_code == 422 and r.json()["error"]["code"] == "no_header"
    big = "Name,Phone,Type\n" + "".join(f"P {i},0300{i:07d},Customer\n" for i in range(2001))
    r = owner.post("/api/v1/bulk/preview/people", **_csv(big))
    assert r.status_code == 422 and r.json()["error"]["code"] == "too_many_rows"
