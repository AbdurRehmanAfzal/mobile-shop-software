"""Owner 2026-10-31: a phone's history and a person's journey download in every format, and every
timeline row says which document it points at."""

import io

from tests.test_purchases import _dev, _party

A = "490154203237518"


def test_phone_history_and_party_journey_download_in_every_format(owner, cat):
    sup = _party(owner, name="Dealer", phone="03005556666", kinds=("supplier",))
    cus = _party(owner, name="Ahmed", phone="03001111111", kinds=("customer",))
    pur = owner.post(
        "/api/v1/purchases",
        json={
            "kind": "supplier",
            "party_id": sup["id"],
            "devices": [{"device": _dev(cat, A), "cost": 800, "asking_price": 1000}],
            "paid": 800,
        },
    ).json()
    dev_id = pur["lines"][0]["device_id"]
    owner.patch(f"/api/v1/devices/{dev_id}", json={"asking_price": 1100, "asking_reason": "market"})
    inv = owner.post(
        "/api/v1/sales",
        json={
            "party_id": cus["id"],
            "devices": [{"device_id": dev_id, "price": 1100}],
            "paid": 600,
        },
    ).json()
    owner.post(
        "/api/v1/payments",
        json={
            "party_id": cus["id"],
            "direction": "in",
            "amount": 500,
            "allocations": [{"document_type": "invoice", "document_id": inv["id"], "amount": 500}],
        },
    )

    # timeline rows point at their documents
    tl = owner.get(f"/api/v1/devices/{dev_id}/money").json()["timeline"]
    by = {r["kind"] if r["kind"] != "event" else r["event"]: r for r in tl}
    assert by["purchase"]["doc_type"] == "purchase" and by["purchase"]["doc_id"] == pur["id"]
    assert by["sale"]["doc_type"] == "invoice" and by["sale"]["doc_id"] == inv["id"]
    assert by["payment_in"]["payment_id"] and by["payment_in"]["doc_no"] == inv["no"]
    assert by["price_changed"]["document_type"] is None and by["price_changed"]["event_id"]

    # the history file: csv, xlsx, printable page
    r = owner.get(f"/api/v1/exports/devices/{dev_id}/history", params={"format": "csv"})
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    text = r.content.decode("utf-8-sig")
    assert "Date,What,Person,Bill,Amount" in text.splitlines()[0]
    assert (
        "price changed" in text and "Ahmed" in text and inv["no"] in text and "payment in" in text
    )
    r = owner.get(f"/api/v1/exports/devices/{dev_id}/history", params={"format": "xlsx"})
    assert r.status_code == 200
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(r.content)).active
    assert [c.value for c in ws[1]][:4] == ["Date", "What", "Person", "Bill"]
    assert ws.max_row >= 5
    r = owner.get(f"/api/v1/exports/devices/{dev_id}/history", params={"format": "pdf"})
    assert r.status_code == 200 and "Phone history" in r.text and A in r.text and "market" in r.text

    # the journey file for Ahmed
    r = owner.get(f"/api/v1/exports/parties/{cus['id']}/journey", params={"format": "csv"})
    assert r.status_code == 200
    text = r.content.decode("utf-8-sig")
    assert (
        "Date,What,No,Phones,Amount" in text.splitlines()[0]
        and inv["no"] in text
        and "payment in" in text
    )
    r = owner.get(f"/api/v1/exports/parties/{cus['id']}/journey", params={"format": "pdf"})
    assert r.status_code == 200 and "Journey" in r.text and "Ahmed" in r.text
