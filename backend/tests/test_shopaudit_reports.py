"""Shop Audit P10 (docs/17): the complete report — same figures as the screens, a sheet per
section in Excel, a printable page for PDF, sections picked by the owner, security only for the owner."""

import io

from tests.conftest import login
from tests.test_devices import _stock
from tests.test_purchases import _party

A, B = "490154203237518", "356938035643809"


def _seed(owner, cat):
    cus = _party(owner, name="Ahmed", phone="03001111111", kinds=("customer",))
    a = _stock(owner, cat, imei=A, condition="box_pack", cost_price=800, asking_price=1000)
    _stock(owner, cat, imei=B, condition="box_pack", cost_price=800, asking_price=1000)
    assert (
        owner.post(
            "/api/v1/sales",
            json={
                "party_id": cus["id"],
                "devices": [{"device_id": a["id"], "price": 1000}],
                "discount": 50,
                "paid": 500,
            },
        ).status_code
        == 201
    )


def test_report_json_xlsx_and_printable_agree_with_the_overview(owner, client, cat):
    _seed(owner, cat)
    q = {"preset": "this_month"}
    j = owner.get("/api/v1/shop-audit/report", params={**q, "format": "json"}).json()
    keys = [s["key"] for s in j["sections"]]
    assert keys[0] == "summary" and "pnl" in keys and "security" in keys and "checks" in keys
    assert j["available"][0] == "summary" and j["titles"]["udhaar"]
    summary = {r[0]: r[1] for r in j["sections"][0]["rows"]}
    o = owner.get("/api/v1/shop-audit/overview", params=q).json()
    assert summary["Net sales"] == o["sales"]["net_sales"] == "950.00"
    assert summary["Net profit"] == o["profit"]["net_profit"] == "150.00"
    assert summary["Cash & bank now"] == o["money"]["current"] == "500.00"
    pnl = {r[0]: r for r in next(s for s in j["sections"] if s["key"] == "pnl")["rows"]}
    assert pnl["Discounts"][1] == "50.00" and pnl["Net sales"][1] == "950.00"
    checks = next(s for s in j["sections"] if s["key"] == "checks")
    assert all(r[3] is True for r in checks["rows"])

    # Excel: a cover sheet plus one per section, headers in row 1, numbers as numbers
    r = owner.get(
        "/api/v1/shop-audit/report",
        params={**q, "format": "xlsx", "sections": "summary,pnl,udhaar"},
    )
    assert r.status_code == 200 and r.headers["content-type"].startswith(
        "application/vnd.openxmlformats"
    )
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(r.content))
    assert (
        wb.sheetnames[0] == "Cover"
        and "Profit & loss" in wb.sheetnames
        and "Executive summary" in wb.sheetnames
    )
    assert "Stock ageing" not in wb.sheetnames
    ws = wb["Profit & loss"]
    assert [c.value for c in ws[1]] == ["Line", "This period", "Period before"]
    net = next(row for row in ws.iter_rows(min_row=2, values_only=True) if row[0] == "Net sales")
    assert net[1] == 950.0
    cover = {row[0]: row[1] for row in wb["Cover"].iter_rows(values_only=True) if row and row[0]}
    assert cover["Currency"] == "PKR" and cover["Mode"] == "summary" and "Period" in cover

    # the printable page: every chosen title, the period line, the detail mode adds the series
    r = owner.get(
        "/api/v1/shop-audit/report",
        params={**q, "format": "pdf", "sections": "summary,pnl,checks", "mode": "detail"},
    )
    assert r.status_code == 200 and "text/html" in r.headers["content-type"]
    import html as _html

    page = _html.unescape(r.text)
    for t in (
        "Shop audit report",
        "Executive summary",
        "Profit & loss",
        "Reconciliation checks",
        "Profit & loss by day",
        "balances as of",
        "PKR",
    ):
        assert t in page, t
    assert "Stock ageing" not in page
    assert owner.get("/api/v1/shop-audit/report", params={"format": "docx"}).status_code == 422

    # a manager gets the report without the security section; a salesman gets nothing
    owner.post(
        "/api/v1/staff",
        json={"name": "Mina", "phone": "03211234568", "password": "secret123", "role": "manager"},
    )
    owner.post(
        "/api/v1/staff",
        json={"name": "Ali", "phone": "03211234567", "password": "secret123", "role": "salesman"},
    )
    mina, ali = login(client, "03211234568"), login(client, "03211234567")
    mj = mina.get("/api/v1/shop-audit/report", params={**q, "format": "json"}).json()
    assert "security" not in [s["key"] for s in mj["sections"]] and "pnl" in [
        s["key"] for s in mj["sections"]
    ]
    assert ali.get("/api/v1/shop-audit/report", params=q).status_code == 403
