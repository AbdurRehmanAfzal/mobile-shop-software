"""FR-PLT-07: CSV exports with BOM, role-gated cost columns, shop isolation."""

from tests.conftest import login
from tests.test_devices import _stock
from tests.test_purchases import _party


def test_stock_and_parties_csv(owner, cat):
    _stock(
        owner, cat, imei="490154203237518", condition="box_pack", cost_price=800, asking_price=1000
    )
    _party(owner, name="احمد خان", phone="03219876543", kinds=("customer",))
    r = owner.get("/api/v1/exports/stock.csv")
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    assert "attachment" in r.headers["content-disposition"]
    body = r.content.decode("utf-8")
    assert body.startswith("﻿") and "490154203237518" in body and "Cost" in body and "800.00" in body
    p = owner.get("/api/v1/exports/parties.csv").content.decode("utf-8")
    assert "احمد خان" in p


def test_salesman_gets_no_cost_and_no_sales_export(client, owner, cat, other_owner):
    _stock(
        owner, cat, imei="490154203237518", condition="box_pack", cost_price=800, asking_price=1000
    )
    owner.post(
        "/api/v1/staff",
        json={"name": "Sales", "phone": "03451112233", "password": "secret123", "role": "salesman"},
    )
    s = login(client, "03451112233")
    body = s.get("/api/v1/exports/stock.csv").content.decode("utf-8")
    assert "Cost" not in body and "800.00" not in body
    assert s.get("/api/v1/exports/sales.csv").status_code == 403
    assert "490154203237518" not in other_owner.get("/api/v1/exports/stock.csv").content.decode(
        "utf-8"
    )


def test_sales_statement_expenses_csv(owner, cat):
    c = _party(owner, name="Ahmed", phone="03219876543", kinds=("customer",))
    d = _stock(
        owner, cat, imei="490154203237518", condition="box_pack", cost_price=800, asking_price=1000
    )
    owner.post(
        "/api/v1/sales",
        json={"party_id": c["id"], "devices": [{"device_id": d["id"], "price": 1000}], "paid": 400},
    )
    s = owner.get("/api/v1/exports/sales.csv").content.decode("utf-8")
    assert "INV-000001" in s and "Ahmed" in s and "200.00" in s  # profit
    st = owner.get(f"/api/v1/exports/parties/{c['id']}/statement.csv").content.decode("utf-8")
    assert "Opening balance" in st and "600.00" in st
    assert owner.get("/api/v1/exports/expenses.csv").status_code == 200
    assert owner.get("/api/v1/exports/purchases.csv").status_code == 200


def test_every_report_downloads_as_csv_excel_and_pdf(owner, cat):
    """Owner: every report needs PDF, Excel and CSV."""
    import io

    from openpyxl import load_workbook

    from tests.test_devices import _stock

    d = _stock(
        owner, cat, imei="490154203237518", condition="box_pack", cost_price=800, asking_price=1000
    )
    owner.post(
        "/api/v1/sales", json={"devices": [{"device_id": d["id"], "price": 1000}], "paid": 1000}
    )
    reports = [
        "stock",
        "parties",
        "sales",
        "purchases",
        "expenses",
        "profit-loss",
        "top-sellers",
        "day",
        "stock-summary",
        "khata",
        "forecast",
        "stock-ageing",
    ]
    for name in reports:
        r = owner.get(f"/api/v1/exports/{name}", params={"format": "csv"})
        assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv"), name
        x = owner.get(f"/api/v1/exports/{name}", params={"format": "xlsx"})
        assert x.status_code == 200 and x.headers["content-disposition"].endswith('.xlsx"'), name
        wb = load_workbook(io.BytesIO(x.content))
        assert wb.active.max_row >= 1, name
        h = owner.get(f"/api/v1/exports/{name}", params={"format": "pdf"})
        assert h.status_code == 200 and h.headers["content-type"].startswith("text/html"), name
        assert "<table>" in h.text and "Test Mobile" in h.text, name
    # numbers stay numbers in Excel
    x = owner.get("/api/v1/exports/profit-loss", params={"format": "xlsx"})
    ws = load_workbook(io.BytesIO(x.content)).active
    assert ws["A2"].value == "Sold" and ws["B2"].value == 1000
    # old .csv paths keep working
    assert owner.get("/api/v1/exports/stock.csv").status_code == 200
